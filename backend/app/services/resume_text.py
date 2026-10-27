"""Text extraction from uploaded resume files (PDF, DOCX, MD, TXT)."""

from __future__ import annotations

import io
import re
from pathlib import Path

SUPPORTED_EXTENSIONS = (".pdf", ".docx", ".md", ".txt")
MAX_FILE_BYTES = 10 * 1024 * 1024


class UnsupportedFormat(ValueError):
    pass


class EmptyText(ValueError):
    pass


class FileTooLarge(ValueError):
    pass


class UnreadableFile(ValueError):
    pass


def normalize_whitespace(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n").replace("\x00", "")
    lines = [re.sub(r"[ \t ]+", " ", line).rstrip() for line in text.split("\n")]
    text = "\n".join(lines)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "utf-8", "cp1251"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def _pdf(data: bytes, *, layout: bool = False) -> str:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    parts: list[str] = []
    for page in reader.pages:
        if layout:
            # layout mode keeps separate text runs apart with runs of spaces (skills of the hh export)
            parts.append(page.extract_text(extraction_mode="layout") or "")
        else:
            parts.append(page.extract_text() or "")
    return "\n\n".join(parts)


def extract_layout_text(data: bytes, filename: str) -> str | None:
    """Layout-preserving text for PDF only (used by the resume parser); None for other formats or on failure."""
    if Path(filename).suffix.lower() != ".pdf":
        return None
    try:
        text = _pdf(data, layout=True)
    except Exception:  # noqa: BLE001 - the parser simply falls back to plain text
        return None
    return text.replace("\r\n", "\n").replace("\x00", "") or None


def _docx(data: bytes) -> str:
    import docx

    document = docx.Document(io.BytesIO(data))
    parts: list[str] = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    return "\n".join(parts)


def extract_text(data: bytes, filename: str) -> str:
    """Returns normalized plain text. Raises UnsupportedFormat / FileTooLarge / UnreadableFile / EmptyText."""
    ext = Path(filename).suffix.lower()
    if ext not in SUPPORTED_EXTENSIONS:
        raise UnsupportedFormat(f"Формат {ext or 'без расширения'} не поддерживается; нужны PDF, DOCX, MD или TXT")
    if len(data) > MAX_FILE_BYTES:
        raise FileTooLarge("Файл больше 10 МБ")
    try:
        if ext == ".pdf":
            text = _pdf(data)
        elif ext == ".docx":
            text = _docx(data)
        else:
            text = _decode(data)
    except Exception as exc:  # noqa: BLE001 - pypdf / python-docx raise many types on a damaged file (was a 500)
        kind = "PDF" if ext == ".pdf" else "DOCX" if ext == ".docx" else "текстовый файл"
        raise UnreadableFile(f"Не удалось прочитать {kind}: файл повреждён или это другой формат") from exc
    text = normalize_whitespace(text)
    if not text:
        raise EmptyText("Из файла не удалось извлечь текст (например, PDF без текстового слоя)")
    return text
