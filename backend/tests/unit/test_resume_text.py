from __future__ import annotations

import io

import pytest

from app.services.resume_text import (
    MAX_FILE_BYTES,
    EmptyText,
    FileTooLarge,
    UnsupportedFormat,
    extract_text,
    normalize_whitespace,
)


def make_pdf(text: str) -> bytes:
    """Minimal single-page PDF with Helvetica text and a correct xref table."""
    content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode("latin-1")
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length " + str(len(content)).encode() + b" >>\nstream\n" + content + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets: list[int] = []
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n".encode() + obj + b"\nendobj\n")
    xref_pos = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n".encode())
    out.write(b"0000000000 65535 f \n")
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref_pos}\n%%EOF\n".encode())
    return out.getvalue()


def make_docx(paragraphs: list[str], table: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for p in paragraphs:
        document.add_paragraph(p)
    if table:
        t = document.add_table(rows=len(table), cols=len(table[0]))
        for r, row in enumerate(table):
            for c, value in enumerate(row):
                t.cell(r, c).text = value
    buf = io.BytesIO()
    document.save(buf)
    return buf.getvalue()


def test_txt_utf8_and_cp1251() -> None:
    assert extract_text("Иванов Иван\n\n\n\nPython".encode(), "cv.txt") == "Иванов Иван\n\nPython"
    assert extract_text("Резюме".encode("cp1251"), "cv.TXT") == "Резюме"
    assert extract_text(b"\xef\xbb\xbf# CV\nline", "cv.md") == "# CV\nline"


def test_pdf_extraction() -> None:
    text = extract_text(make_pdf("Hello PDF resume"), "cv.pdf")
    assert "Hello PDF resume" in text


def test_docx_extraction_with_table() -> None:
    data = make_docx(["Иванов Иван", "Python-разработчик"], [["Навык", "Уровень"], ["Python", "5"]])
    text = extract_text(data, "cv.docx")
    assert "Иванов Иван" in text and "Python-разработчик" in text
    assert "Python | 5" in text


def test_unsupported_and_limits() -> None:
    with pytest.raises(UnsupportedFormat):
        extract_text(b"x", "cv.rtf")
    with pytest.raises(FileTooLarge):
        extract_text(b"x" * (MAX_FILE_BYTES + 1), "cv.txt")
    with pytest.raises(EmptyText):
        extract_text(b"   \n\n  ", "cv.txt")


def test_normalize_whitespace() -> None:
    assert normalize_whitespace("a  \t b\r\n\r\n\r\n\r\nc \x00") == "a b\n\nc"
